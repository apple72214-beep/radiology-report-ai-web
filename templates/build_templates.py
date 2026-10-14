#!/usr/bin/env python3
"""Reproducible builder for the docxtpl master templates (EN + AR).

    pip install python-docx
    python3 templates/build_templates.py            # rebuild both .docx
    python3 templates/build_templates.py --check    # rebuild in memory, diff vs on-disk

CRITICAL docxtpl syntax (learned the hard way, 2026-09-20):
  RichText placeholders MUST be written {{r var}} — letter r ADJACENT to the opening
  braces, no space. "{{ r var }}" is invalid Jinja (error "got 'r'"); "{{ var|r }}"
  renders SILENTLY EMPTY because only the {{r prefix triggers docxtpl patch_xml run
  splitting. Never "fix" these tags in the builder below.
"""
import io
import os
import sys
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Mm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
DBLUE = RGBColor(0x00, 0x00, 0x8B)
GRAY = RGBColor(0x80, 0x80, 0x80)
BOX = "EBEBEB"
HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "org_logo.png")

# ---- DRAFT letterhead: swap these constants (and org_logo.png) with the real
# ---- hospital identity, then re-run this builder + backend/render_master_template.py
ORG = {
    "name_en": "City Teaching Hospital",
    "dept_en": "Department of Radiology & Medical Imaging",
    "addr_en": "Street 12 • PO Box 144 • Sana'a • Tel +967-1-000000",
    "name_ar": "المستشفى التعليمي بالمدينة",
    "dept_ar": "قسم الأشعة والتصوير الطبي",
    "addr_ar": "شارع 12 • ص.ب 144 • صنعاء • هاتف 000000 1 967+",
}


def setup(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(15)
    s.left_margin = s.right_margin = Mm(15)


def bidi(p):
    p._p.get_or_add_pPr().append(parse_xml(r'<w:bidi %s/>' % nsdecls('w')))


def shade(p, fill=BOX):
    p._p.get_or_add_pPr().append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), fill)))


def run(p, text, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = size
    return r


def letterhead(doc, lang):
    """Two-cell draft letterhead: logo + org block, then a navy rule."""
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = False
    c0, c1 = tbl.rows[0].cells
    logo_cell, txt_cell = (c0, c1) if lang == "en" else (c1, c0)
    logo_cell.width, txt_cell.width = Mm(22), Mm(168)
    lp = logo_cell.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.add_run().add_picture(LOGO, width=Mm(17))
    lines = [("name", True, NAVY, Pt(13)), ("dept", True, DBLUE, Pt(10)), ("addr", False, GRAY, Pt(8))]
    first = True
    for key, bold, color, size in lines:
        tp = txt_cell.paragraphs[0] if first else txt_cell.add_paragraph()
        first = False
        tp.paragraph_format.space_after = Pt(1)
        if lang == "ar":
            bidi(tp)
            tp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run(tp, ORG[key + "_" + lang], bold=bold, color=color, size=size)
    rule = doc.add_paragraph()
    rule._p.get_or_add_pPr().append(parse_xml(
        r'<w:pBdr %s><w:bottom w:val="single" w:sz="8" w:space="1" w:color="1F3864"/></w:pBdr>' % nsdecls('w')))
    rule.paragraph_format.space_after = Pt(8)


def build_en():
    doc = Document()
    setup(doc)
    letterhead(doc, "en")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(t, "MEDICAL RADIOLOGY REPORT", bold=True, color=NAVY, size=Pt(15))
    for line in ["Patient Name: {{patient_name}}",
                 "Sex: {{sex}}     Age: {{age}}",
                 "Exam Date: {{exam_date}}",
                 "Exam: {{exam}}",
                 "Clinical Indication: {{clinical_indication}}."]:
        doc.add_paragraph(line)
    h = doc.add_paragraph(); run(h, "TECHNIQUE:", bold=True, color=DBLUE)
    doc.add_paragraph("{{technique}}")
    h = doc.add_paragraph(); run(h, "FINDINGS:", bold=True, color=DBLUE)
    doc.add_paragraph("{%p for f in findings_list %}")
    doc.add_paragraph("{{r f}}")
    doc.add_paragraph("{%p endfor %}")
    h = doc.add_paragraph(); run(h, "IMPRESSION:", bold=True, color=DBLUE)
    doc.add_paragraph("{%p for i in impression_list %}")
    p = doc.add_paragraph(); shade(p); p.paragraph_format.space_after = Pt(0)
    run(p, "• "); run(p, "{{r i}}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("")
    sig = doc.add_paragraph(); sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(sig, "Reported by:\n{{radiologist_name}}\n{{title}}")
    foot = doc.add_paragraph()
    run(foot, "Document SHA-256: {{sha256}} — Signed at: {{signed_at}} — Decision-support output; final clinical responsibility rests with the signing physician.", color=GRAY, size=Pt(8))
    return doc


def build_ar():
    doc = Document()
    setup(doc)
    letterhead(doc, "ar")
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; bidi(t)
    run(t, "تقرير أشعة طبي", bold=True, color=NAVY, size=Pt(15))
    for line in ["اسم المريض: {{patient_name}}",
                 "الجنس: {{sex}}     العمر: {{age}}",
                 "تاريخ الفحص: {{exam_date}}",
                 "الفحص: {{exam}}",
                 "الدلالة السريرية: {{clinical_indication}}."]:
        p = doc.add_paragraph(line); bidi(p)
    p = doc.add_paragraph(); bidi(p); run(p, "التقنية:", bold=True, color=DBLUE)
    p = doc.add_paragraph("{{technique}}"); bidi(p)
    p = doc.add_paragraph(); bidi(p); run(p, "الوصف:", bold=True, color=DBLUE)
    p = doc.add_paragraph("{%p for f in findings_list %}"); bidi(p)
    p = doc.add_paragraph("{{r f}}"); bidi(p)
    p = doc.add_paragraph("{%p endfor %}"); bidi(p)
    p = doc.add_paragraph(); bidi(p); run(p, "الانطباع:", bold=True, color=DBLUE)
    p = doc.add_paragraph("{%p for i in impression_list %}"); bidi(p)
    p = doc.add_paragraph(); bidi(p); shade(p); p.paragraph_format.space_after = Pt(0)
    run(p, "• "); run(p, "{{r i}}")
    p = doc.add_paragraph("{%p endfor %}"); bidi(p)
    p = doc.add_paragraph(""); bidi(p)
    sig = doc.add_paragraph(); sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT; bidi(sig)
    run(sig, "التقرير بواسطة:\n{{radiologist_name}}\n{{title}}")
    p = doc.add_paragraph(); bidi(p)
    run(p, "بصمة المستند SHA-256: {{sha256}} — حُرر في: {{signed_at}} — مخرج لدعم القرار؛ المسؤولية السريرية النهائية على الطبيب الموقّع.", color=GRAY, size=Pt(8))
    return doc


def xml_of(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    with zipfile.ZipFile(buf) as z:
        return z.read("word/document.xml")


def main():
    targets = {"radiology_master_template.docx": build_en(), "radiology_master_template_AR.docx": build_ar()}
    if "--check" in sys.argv:
        ok = True
        for name, doc in targets.items():
            same = xml_of(doc) == open(os.path.join(HERE, name), "rb").read().split(b"word/document.xml")[-1] if False else None
            disk = zipfile.ZipFile(os.path.join(HERE, name)).read("word/document.xml")
            ok &= xml_of(doc) == disk
            print(name, "IDENTICAL" if xml_of(doc) == disk else "DIFFERS")
        sys.exit(0 if ok else 1)
    for name, doc in targets.items():
        doc.save(os.path.join(HERE, name))
        print("built", name)


if __name__ == "__main__":
    main()
